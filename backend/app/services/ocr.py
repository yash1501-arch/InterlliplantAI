import io
import math
from dataclasses import dataclass, field
from typing import Any

import openpyxl
from PIL import Image
import pytesseract
from docx import Document
from pypdf import PdfReader


@dataclass
class OCRResult:
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)


def _build_common_metadata(filename: str, file_type: str, text: str) -> dict[str, Any]:
    char_count = len(text)
    word_count = len(text.split()) if text.strip() else 0
    estimated_page_count = max(1, math.ceil(char_count / 3000))
    return {
        "filename": filename,
        "file_type": file_type,
        "character_count": char_count,
        "char_count": char_count,
        "word_count": word_count,
        "estimated_page_count": estimated_page_count,
    }


def _extract_txt(filename: str, content: bytes) -> OCRResult:
    decoded_text = content.decode("utf-8", errors="ignore")
    lines = [line.strip() for line in decoded_text.splitlines() if line.strip()]
    text = "\n".join(lines)

    sections = []
    for line in lines:
        stripped = line.strip()
        if stripped.endswith(":"):
            sections.append(stripped[:-1].strip())
        elif ":" in stripped:
            heading = stripped.split(":", 1)[0].strip()
            if heading:
                sections.append(heading)

    keywords = [
        token.strip(".,;:-()[]{}\"")
        for token in text.split()
        if len(token.strip(".,;:-()[]{}\"") or "") > 3
    ]
    unique_keywords = sorted(set(keywords))[:12]

    metadata = _build_common_metadata(filename, "txt", text)
    metadata["line_count"] = len(lines)
    metadata["section_count"] = len(sections)
    metadata["keywords"] = unique_keywords
    metadata["source"] = "starter-ocr-pipeline"

    return OCRResult(text=text, metadata=metadata)


def _extract_pdf(filename: str, content: bytes) -> OCRResult:
    try:
        reader = PdfReader(io.BytesIO(content))
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
        text = "\n".join(text_parts)
        metadata = _build_common_metadata(filename, "pdf", text)
        metadata["page_count"] = len(reader.pages)
        metadata["estimated_page_count"] = len(reader.pages)
        return OCRResult(text=text, metadata=metadata)
    except Exception:
        metadata = _build_common_metadata(filename, "pdf", "")
        metadata["page_count"] = 0
        metadata["extraction_error"] = "Failed to extract text from PDF"
        return OCRResult(text="", metadata=metadata)


def _extract_docx(filename: str, content: bytes) -> OCRResult:
    try:
        doc = Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs]
        text = "\n".join(paragraphs)
        metadata = _build_common_metadata(filename, "docx", text)
        return OCRResult(text=text, metadata=metadata)
    except Exception:
        metadata = _build_common_metadata(filename, "docx", "")
        metadata["extraction_error"] = "Failed to extract text from DOCX"
        return OCRResult(text="", metadata=metadata)


def _extract_xlsx(filename: str, content: bytes) -> OCRResult:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        text_parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    text_parts.append(row_text)
        text = "\n".join(text_parts)
        metadata = _build_common_metadata(filename, "xlsx", text)
        return OCRResult(text=text, metadata=metadata)
    except Exception:
        metadata = _build_common_metadata(filename, "xlsx", "")
        metadata["extraction_error"] = "Failed to extract text from XLSX"
        return OCRResult(text="", metadata=metadata)


def _extract_csv(filename: str, content: bytes) -> OCRResult:
    try:
        decoded = content.decode("utf-8", errors="ignore")
        import csv
        reader = csv.reader(io.StringIO(decoded))
        rows = []
        for row in reader:
            rows.append(",".join(cell.strip() for cell in row))
        text = "\n".join(rows)
        metadata = _build_common_metadata(filename, "csv", text)
        return OCRResult(text=text, metadata=metadata)
    except Exception:
        metadata = _build_common_metadata(filename, "csv", "")
        metadata["extraction_error"] = "Failed to parse CSV"
        return OCRResult(text="", metadata=metadata)


def _extract_image(filename: str, content: bytes) -> OCRResult:
    try:
        image = Image.open(io.BytesIO(content))
        text = pytesseract.image_to_string(image)
        metadata = _build_common_metadata(filename, filename.rsplit(".", 1)[-1].lower(), text)
        metadata["estimated_page_count"] = 1
        return OCRResult(text=text, metadata=metadata)
    except Exception:
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "image"
        metadata = _build_common_metadata(filename, ext, "")
        metadata["estimated_page_count"] = 1
        metadata["extraction_error"] = "OCR failed on image"
        return OCRResult(text="", metadata=metadata)


_IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "tiff", "tif", "bmp"}


def extract_text_and_metadata(filename: str, content: bytes) -> OCRResult:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "txt":
        return _extract_txt(filename, content)
    elif ext == "pdf":
        return _extract_pdf(filename, content)
    elif ext == "docx":
        return _extract_docx(filename, content)
    elif ext == "xlsx":
        return _extract_xlsx(filename, content)
    elif ext == "csv":
        return _extract_csv(filename, content)
    elif ext in _IMAGE_EXTENSIONS:
        return _extract_image(filename, content)
    else:
        raise ValueError(f"Unsupported file type: .{ext}")
